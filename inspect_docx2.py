from docx import Document
from docx.shared import Pt, RGBColor
import re

doc = Document('KP4.docx')

print('=== PARAGRAPH DETAILED ANALYSIS ===')
for i, para in enumerate(doc.paragraphs):
    if para.text.strip():
        print('\nPara %d:' % i)
        print('  Style: %s' % para.style.name)
        print('  Text: "%s"' % para.text)
        print('  Alignment: %s' % para.alignment)
        for j, run in enumerate(para.runs):
            if run.text.strip():
                print('  Run %d: "%s" | Bold=%s | Italic=%s | Size=%s | Font=%s' % (
                    j, run.text, run.bold, run.italic, 
                    run.font.size.pt if run.font.size else 'inherited',
                    run.font.name
                ))

print('\n=== TABLE DETAILED ANALYSIS ===')
for t_idx, table in enumerate(doc.tables):
    print('\n--- Table %d ---' % t_idx)
    print('  Style: %s' % table.style.name if table.style else 'None')
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            if cell.text.strip():
                print('  Cell[%d][%d]: "%s"' % (r_idx, c_idx, cell.text.strip()))
                for para in cell.paragraphs:
                    for run in para.runs:
                        if run.text.strip():
                            print('    Run: "%s" | Bold=%s | Size=%s | Font=%s' % (
                                run.text, run.bold,
                                run.font.size.pt if run.font.size else 'inherited',
                                run.font.name
                            ))

print('\n=== DOCUMENT SECTIONS/MARGINS ===')
for section in doc.sections:
    print('  Page width: %s cm' % (section.page_width.cm))
    print('  Page height: %s cm' % (section.page_height.cm))
    print('  Left margin: %s cm' % (section.left_margin.cm))
    print('  Right margin: %s cm' % (section.right_margin.cm))
    print('  Top margin: %s cm' % (section.top_margin.cm))
    print('  Bottom margin: %s cm' % (section.bottom_margin.cm))
