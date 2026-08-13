from docx import Document

doc = Document('KP4.docx')

print('=== DOCUMENT STRUCTURE ===')
print('Total paragraphs:', len(doc.paragraphs))
print('Total tables:', len(doc.tables))
print()

print('=== ALL PARAGRAPHS ===')
for i, para in enumerate(doc.paragraphs):
    if para.text.strip():
        print('Para %d: [%s] "%s"' % (i, para.style.name, para.text))

print()
print('=== ALL TABLES ===')
for t_idx, table in enumerate(doc.tables):
    print('\n--- Table %d (%d rows x %d cols) ---' % (t_idx, len(table.rows), len(table.columns)))
    for r_idx, row in enumerate(table.rows):
        row_texts = []
        for c_idx, cell in enumerate(row.cells):
            text = cell.text.strip()
            row_texts.append('[%d]:%s' % (c_idx, text))
        print('  Row %d: %s' % (r_idx, ' | '.join(row_texts)))
